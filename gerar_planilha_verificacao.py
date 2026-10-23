"""
Gera uma planilha do Excel que confere os cálculos do simulador.

O que a planilha faz:
- refaz, com FÓRMULAS do Excel, todas as contas do simulador (chance, total
  apostado, perda esperada, investimento);
- reproduz a simulação nº 1 período por período, com os mesmos sorteios
  que o Python usou;
- refaz as estatísticas (resultado típico, faixa de 90%, chance de lucro) a
  partir dos resultados das 2.000 simulações;
- compara tudo com o que o Python calculou e marca "OK" ou "VERIFICAR".

Como usar (no terminal, dentro da pasta do projeto):
    python gerar_planilha_verificacao.py
Para conferir outro cenário, mude os valores em PARAMETROS abaixo e rode de novo.
"""

from __future__ import annotations

import numpy as np
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

import config
from simulador import PERIODOS_POR_ANO, Parametros, _converter_frequencia, simular

# ---------------------------------------------------------------------------
# Cenário a conferir (os mesmos campos da tela do simulador)
# ---------------------------------------------------------------------------
PARAMETROS = Parametros(
    valor_aposta=20.0,        # Valor de cada aposta (R$)
    quantidade=3,             # Quantas apostas
    frequencia="semana",      # "dia", "semana" ou "mês"
    anos=3,                   # Por quantos anos
    odd_media=2.0,            # Odd média das apostas
    margem_casa=0.07,         # Margem da casa (0,07 = 7%)
    taxa_investimento=0.10,   # Rendimento do investimento (0,10 = 10% ao ano)
)
ARQUIVO = "verificacao_simulador.xlsx"

# ---------------------------------------------------------------------------
# Estilo
# ---------------------------------------------------------------------------
FONTE = "Arial"
F_NORMAL = Font(name=FONTE, size=10)
F_NEGRITO = Font(name=FONTE, size=10, bold=True)
F_TITULO = Font(name=FONTE, size=14, bold=True)
F_ENTRADA = Font(name=FONTE, size=10, color="0000FF")       # azul = dado de entrada
F_PYTHON = Font(name=FONTE, size=10, color="008000")        # verde = valor vindo do Python
F_NOTA = Font(name=FONTE, size=9, italic=True, color="595959")
FUNDO_CAB = PatternFill("solid", fgColor="DCE6F2")
FUNDO_ENTRADA = PatternFill("solid", fgColor="FFF2CC")
FUNDO_OK = PatternFill("solid", fgColor="E2EFDA")
FUNDO_ERRO = PatternFill("solid", fgColor="F8CBAD")
LINHA = Side(style="thin", color="BFBFBF")
BORDA = Border(left=LINHA, right=LINHA, top=LINHA, bottom=LINHA)
BRL = 'R$ #,##0.00;-R$ #,##0.00;R$ 0.00'
PCT = "0.00%"
TOLERANCIA = 0.01  # diferença aceitável entre Excel e Python (1 centavo)


def cabecalho(ws, linha, textos, larguras=None):
    for j, t in enumerate(textos, start=1):
        c = ws.cell(linha, j, t)
        c.font, c.fill, c.border = F_NEGRITO, FUNDO_CAB, BORDA
        c.alignment = Alignment(wrap_text=True, vertical="center")
    if larguras:
        for j, w in enumerate(larguras, start=1):
            ws.column_dimensions[chr(64 + j)].width = w


def celula(ws, ref, valor, fonte=F_NORMAL, formato=None, fundo=None):
    c = ws[ref]
    c.value, c.font, c.border = valor, fonte, BORDA
    if formato:
        c.number_format = formato
    if fundo:
        c.fill = fundo
    return c


def colorir_status(ws, intervalo):
    ws.conditional_formatting.add(intervalo, CellIsRule(operator="equal", formula=['"OK"'], fill=FUNDO_OK))
    ws.conditional_formatting.add(intervalo, CellIsRule(operator="equal", formula=['"VERIFICAR"'], fill=FUNDO_ERRO))


def status(dif_ref):
    return f'=IF(ABS({dif_ref})<={TOLERANCIA},"OK","VERIFICAR")'


# ---------------------------------------------------------------------------
# 1. Rodar o simulador e recriar os mesmos sorteios
# ---------------------------------------------------------------------------
p = PARAMETROS
r = simular(p)
s = r.resumo
periodo, apostas_por_periodo = _converter_frequencia(p)
n_periodos = p.anos * PERIODOS_POR_ANO[periodo]
chance = (1 / p.odd_media) * (1 - p.margem_casa)
# Mesma semente e mesma chamada do simulador.py -> exatamente os mesmos números
vitorias = np.random.default_rng(p.semente).binomial(
    apostas_por_periodo, chance, size=(p.simulacoes, n_periodos))

wb = Workbook()

# ---------------------------------------------------------------------------
# Aba "Como usar"
# ---------------------------------------------------------------------------
ws = wb.active
ws.title = "Como usar"
texto = [
    ("Verificação dos cálculos do simulador da Betinha", F_TITULO),
    ("", None),
    ("O que é: esta planilha refaz, com fórmulas do Excel, as contas que o simulador faz em Python, "
     "e compara os dois resultados.", None),
    ("Como ler: na coluna \"Situação\", OK (verde) quer dizer que o Excel e o Python chegaram ao mesmo número "
     f"(diferença de até R$ {TOLERANCIA:.2f}). VERIFICAR (vermelho) indica que algo precisa ser revisado.", None),
    ("", None),
    ("Cores:", F_NEGRITO),
    ("  Azul com fundo amarelo = dados de entrada (os mesmos da tela do simulador).", None),
    ("  Preto = fórmula do Excel.", None),
    ("  Verde = número calculado pelo Python (simulador.py), copiado para comparação.", None),
    ("", None),
    ("Abas:", F_NEGRITO),
    ("  1. Resumo – quantas verificações deram OK.", None),
    ("  2. Fórmulas – conta por conta: chance real, total apostado, perda esperada, investimento.", None),
    ("  3. Simulação 1 – a primeira das 2.000 simulações, período por período, com os mesmos sorteios do Python.", None),
    ("  4. Investimento – o dinheiro investido crescendo período por período (juros compostos).", None),
    ("  5. 2.000 simulações – o resultado final de cada simulação e as estatísticas mostradas na tela.", None),
    ("", None),
    ("Para conferir outro cenário: mude PARAMETROS no arquivo gerar_planilha_verificacao.py e rode "
     "\"python gerar_planilha_verificacao.py\" de novo.", None),
    ("Observação: a média das 2.000 simulações é aleatória. Ela não precisa ser igual à perda esperada, só "
     "ficar perto dela (Lei dos Grandes Números). A aba Fórmulas mostra a margem de tolerância usada.", None),
]
for i, (t, f) in enumerate(texto, start=1):
    c = ws.cell(i, 1, t)
    c.font = f or F_NORMAL
    c.alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 115

# ---------------------------------------------------------------------------
# Aba "Fórmulas"
# ---------------------------------------------------------------------------
wf = wb.create_sheet("Fórmulas")
wf["A1"] = "Conta por conta: Excel × Python"
wf["A1"].font = F_TITULO
cabecalho(wf, 3, ["Entrada (igual à tela do simulador)", "Valor"], [46, 18, 20, 14, 13, 55])
entradas = [
    ("Valor de cada aposta (R$)", p.valor_aposta, BRL),
    ("Quantas apostas", p.quantidade, "0"),
    ("Com que frequência (dia, semana ou mês)", p.frequencia, "@"),
    ("Por quantos anos", p.anos, "0"),
    ("Odd média das apostas", p.odd_media, "0.00"),
    ("Margem da casa", p.margem_casa, PCT),
    ("Rendimento do investimento (ao ano)", p.taxa_investimento, PCT),
]
for k, (nome, valor, fmt) in enumerate(entradas, start=4):
    celula(wf, f"A{k}", nome)
    celula(wf, f"B{k}", valor, F_ENTRADA, fmt, FUNDO_ENTRADA)
dv = DataValidation(type="list", formula1='"dia,semana,mês"', allow_blank=False)
wf.add_data_validation(dv)
dv.add("B6")
# B4 valor · B5 quantidade · B6 frequência · B7 anos · B8 odd · B9 margem · B10 taxa

cabecalho(wf, 12, ["Conta", "Excel (fórmula)", "Python (simulador)", "Diferença", "Situação", "Como é calculado"])
contas = [
    ("Chance que a odd mostra", "=1/B8", s["chance_implicita"], PCT, "1 ÷ odd"),
    ("Chance real de ganhar", "=(1/B8)*(1-B9)", s["chance_real"], PCT, "(1 ÷ odd) × (1 − margem)"),
    ("Apostas por período", '=IF(B6="dia",B5*7,B5)', apostas_por_periodo, "0",
     "se for por dia, o simulador agrupa por semana (× 7)"),
    ("Períodos por ano", '=IF(B6="mês",12,52)', PERIODOS_POR_ANO[periodo], "0", "52 semanas ou 12 meses"),
    ("Número de apostas no total", "=B15*B16*B7", s["numero_de_apostas"], "#,##0",
     "apostas por período × períodos por ano × anos"),
    ("Total apostado (R$)", "=B17*B4", s["total_apostado"], BRL, "número de apostas × valor"),
    ("Ganho ou perda média por aposta (R$)", "=B4*(B14*B8-1)", -p.margem_casa * p.valor_aposta, BRL,
     "valor × (chance real × odd − 1)"),
    ("Perda esperada no período (R$)", "=B18*B9", s["perda_esperada"], BRL, "margem × total apostado"),
    ("Mesma perda, pela média por aposta (R$)", "=-B19*B17", s["perda_esperada"], BRL,
     "deve dar igual à linha de cima"),
    ("Taxa do investimento por período", "=(1+B10)^(1/B16)-1", (1 + p.taxa_investimento) ** (1 / PERIODOS_POR_ANO[periodo]) - 1,
     "0.0000%", "(1 + taxa anual) ^ (1 ÷ períodos por ano) − 1"),
    ("Valor guardado por período (R$)", "=B15*B4", apostas_por_periodo * p.valor_aposta, BRL,
     "o mesmo dinheiro que iria para as apostas"),
    ("Saldo se tivesse investido (R$)", "=IF(B22=0,B23*B16*B7,B23*(((1+B22)^(B16*B7)-1)/B22)*(1+B22))",
     s["saldo_investindo"], BRL, "aporte × [((1 + i)^n − 1) ÷ i] × (1 + i)"),
    ("Rendimento do investimento (R$)", "=B24-B18", s["rendimento_investindo"], BRL, "saldo − total guardado"),
]
# "O que esse dinheiro compraria" (preços de referência do config.py)
for e in config.EQUIVALENCIAS:
    contas.append((f"Perda média em {e['plural']}", f"=B20/{e['valor']}", s["perda_esperada"] / e["valor"],
                   "0.00", f"perda esperada ÷ R$ {e['valor']:.2f} ({e['fonte']})"))
for k, (nome, formula, py, fmt, como) in enumerate(contas, start=13):
    celula(wf, f"A{k}", nome)
    celula(wf, f"B{k}", formula, formato=fmt)
    celula(wf, f"C{k}", float(py), F_PYTHON, fmt)
    celula(wf, f"D{k}", f"=B{k}-C{k}", formato="0.0000")
    celula(wf, f"E{k}", status(f"D{k}"))
    celula(wf, f"F{k}", como, F_NOTA)
ultima_conta = 12 + len(contas)
colorir_status(wf, f"E13:E{ultima_conta}")

# Média do Monte Carlo × perda esperada (aleatório: compara com tolerância estatística)
k = ultima_conta + 2
wf[f"A{k}"] = "Média das 2.000 simulações × perda esperada"
wf[f"A{k}"].font = F_NEGRITO
celula(wf, f"A{k+1}", "Resultado médio das simulações (R$)")
celula(wf, f"B{k+1}", "=AVERAGE('2.000 simulações'!B:B)", formato=BRL)
celula(wf, f"C{k+1}", float(s["resultado_medio_simulado"]), F_PYTHON, BRL)
celula(wf, f"A{k+2}", "Perda esperada pela fórmula (R$)")
celula(wf, f"B{k+2}", "=-B20", formato=BRL)
celula(wf, f"A{k+3}", "Tolerância (3 erros-padrão) (R$)")
celula(wf, f"B{k+3}", "=3*STDEV('2.000 simulações'!B:B)/SQRT(COUNT('2.000 simulações'!B:B))", formato=BRL)
celula(wf, f"A{k+4}", "A média ficou perto do esperado?")
celula(wf, f"B{k+4}", f'=IF(ABS(B{k+1}-B{k+2})<=B{k+3},"OK","VERIFICAR")')
colorir_status(wf, f"B{k+4}")
wf[f"F{k+1}"] = "É sorteio: não precisa dar igual, só ficar dentro da tolerância."
wf[f"F{k+1}"].font = F_NOTA
linha_media_ok = k + 4

# ---------------------------------------------------------------------------
# Aba "Simulação 1": a primeira trajetória, período por período
# ---------------------------------------------------------------------------
w1 = wb.create_sheet("Simulação 1")
w1["A1"] = "A simulação nº 1, período por período (mesmos sorteios do Python)"
w1["A1"].font = F_TITULO
w1["A2"] = ("A coluna B vem do Python: quantas apostas a pessoa ganhou em cada período. "
            "Todo o resto é fórmula do Excel.")
w1["A2"].font = F_NOTA
cabecalho(w1, 4, ["Período", "Apostas ganhas (sorteio do Python)", "Apostas feitas", "Recebido (R$)",
                  "Gasto (R$)", "Lucro do período (R$)", "Lucro acumulado (R$)"],
          [10, 18, 12, 15, 14, 18, 20])
for t in range(n_periodos):
    lin = 5 + t
    celula(w1, f"A{lin}", t + 1, formato="0")
    celula(w1, f"B{lin}", int(vitorias[0, t]), F_PYTHON, "0")
    celula(w1, f"C{lin}", "='Fórmulas'!$B$15", formato="0")
    celula(w1, f"D{lin}", f"=B{lin}*'Fórmulas'!$B$4*'Fórmulas'!$B$8", formato=BRL)
    celula(w1, f"E{lin}", f"=C{lin}*'Fórmulas'!$B$4", formato=BRL)
    celula(w1, f"F{lin}", f"=D{lin}-E{lin}", formato=BRL)
    celula(w1, f"G{lin}", f"=F{lin}" if t == 0 else f"=G{lin-1}+F{lin}", formato=BRL)
fim1 = 4 + n_periodos
celula(w1, "I4", "Conferência", F_NEGRITO, fundo=FUNDO_CAB)
celula(w1, "J4", "", fundo=FUNDO_CAB)
celula(w1, "I5", "Resultado final no Excel (R$)")
celula(w1, "J5", f"=G{fim1}", formato=BRL)
celula(w1, "I6", "Resultado final no Python (R$)")
celula(w1, "J6", float(r.resultados_finais[0]), F_PYTHON, BRL)
celula(w1, "I7", "Situação")
celula(w1, "J7", status("J5-J6"))
colorir_status(w1, "J7")
w1.column_dimensions["I"].width = 30
w1.column_dimensions["J"].width = 16
w1.freeze_panes = "A5"

# ---------------------------------------------------------------------------
# Aba "Investimento": juros compostos período por período
# ---------------------------------------------------------------------------
wi = wb.create_sheet("Investimento")
wi["A1"] = "Se o mesmo dinheiro fosse investido (aplicado no início de cada período)"
wi["A1"].font = F_TITULO
cabecalho(wi, 3, ["Período", "Saldo no início (R$)", "Depósito (R$)", "Rendimento do período (R$)",
                  "Saldo no fim (R$)"], [10, 20, 14, 24, 18])
for t in range(n_periodos):
    lin = 4 + t
    celula(wi, f"A{lin}", t + 1, formato="0")
    celula(wi, f"B{lin}", 0 if t == 0 else f"=E{lin-1}", formato=BRL)
    celula(wi, f"C{lin}", "='Fórmulas'!$B$23", formato=BRL)
    celula(wi, f"D{lin}", f"=(B{lin}+C{lin})*'Fórmulas'!$B$22", formato=BRL)
    celula(wi, f"E{lin}", f"=B{lin}+C{lin}+D{lin}", formato=BRL)
fimi = 3 + n_periodos
celula(wi, "G3", "Conferência", F_NEGRITO, fundo=FUNDO_CAB)
celula(wi, "H3", "", fundo=FUNDO_CAB)
celula(wi, "G4", "Saldo final, período a período (R$)")
celula(wi, "H4", f"=E{fimi}", formato=BRL)
celula(wi, "G5", "Saldo final no Python (R$)")
celula(wi, "H5", float(s["saldo_investindo"]), F_PYTHON, BRL)
celula(wi, "G6", "Situação")
celula(wi, "H6", status("H4-H5"))
colorir_status(wi, "H6")
wi.column_dimensions["G"].width = 34
wi.column_dimensions["H"].width = 16
wi.freeze_panes = "A4"

# ---------------------------------------------------------------------------
# Aba "2.000 simulações": estatísticas mostradas na tela
# ---------------------------------------------------------------------------
wm = wb.create_sheet("2.000 simulações")
cabecalho(wm, 1, ["Simulação", "Lucro ou prejuízo final (R$)"], [12, 26, 4, 46, 18, 20, 14, 13])
for i, v in enumerate(r.resultados_finais):
    celula(wm, f"A{i+2}", i + 1, formato="0")
    celula(wm, f"B{i+2}", float(v), F_PYTHON, BRL)
fimm = 1 + len(r.resultados_finais)
faixa = f"B2:B{fimm}"
cabecalho_est = ["Número que aparece na tela", "Excel (fórmula)", "Python (simulador)", "Diferença", "Situação"]
for j, t in enumerate(cabecalho_est):
    c = wm.cell(1, 4 + j, t)
    c.font, c.fill, c.border = F_NEGRITO, FUNDO_CAB, BORDA
estat = [
    ("Resultado típico (mediana)", f"=MEDIAN({faixa})", s["resultado_mediano"], BRL),
    ("Pior caso da faixa (5%)", f"=PERCENTILE({faixa},0.05)", s["pior_cenario_5"], BRL),
    ("Melhor caso da faixa (95%)", f"=PERCENTILE({faixa},0.95)", s["melhor_cenario_95"], BRL),
    ("Chance de terminar no lucro", f'=COUNTIF({faixa},">0")/COUNT({faixa})', s["chance_terminar_no_lucro"], PCT),
    ("Resultado médio", f"=AVERAGE({faixa})", s["resultado_medio_simulado"], BRL),
]
for k, (nome, formula, py, fmt) in enumerate(estat, start=2):
    celula(wm, f"D{k}", nome)
    celula(wm, f"E{k}", formula, formato=fmt)
    celula(wm, f"F{k}", float(py), F_PYTHON, fmt)
    celula(wm, f"G{k}", f"=E{k}-F{k}", formato="0.0000")
    celula(wm, f"H{k}", status(f"G{k}"))
colorir_status(wm, f"H2:H{1 + len(estat)}")
wm.freeze_panes = "A2"

# ---------------------------------------------------------------------------
# Aba "Resumo" (fica em 2º lugar)
# ---------------------------------------------------------------------------
wr = wb.create_sheet("Resumo", 1)
wr["A1"] = "Resumo da verificação"
wr["A1"].font = F_TITULO
cabecalho(wr, 3, ["Verificação", "Situação"], [60, 16])
checks = [
    ("Contas da aba Fórmulas (todas)",
     f"=IF(COUNTIF('Fórmulas'!E13:E{ultima_conta},\"OK\")={len(contas)},\"OK\",\"VERIFICAR\")"),
    ("Média das simulações perto da perda esperada", f"='Fórmulas'!B{linha_media_ok}"),
    ("Simulação nº 1 refeita período por período", "='Simulação 1'!J7"),
    ("Investimento refeito período por período", "=Investimento!H6"),
    ("Estatísticas das 2.000 simulações (todas)",
     f"=IF(COUNTIF('2.000 simulações'!H2:H{1 + len(estat)},\"OK\")={len(estat)},\"OK\",\"VERIFICAR\")"),
]
for k, (nome, formula) in enumerate(checks, start=4):
    celula(wr, f"A{k}", nome)
    celula(wr, f"B{k}", formula)
fimr = 3 + len(checks)
colorir_status(wr, f"B4:B{fimr}")
celula(wr, f"A{fimr + 2}", "Verificações OK", F_NEGRITO)
celula(wr, f"B{fimr + 2}", f'=COUNTIF(B4:B{fimr},"OK")&" de {len(checks)}"', F_NEGRITO)
wr[f"A{fimr + 4}"] = (f"Cenário conferido: {p.quantidade} aposta(s) por {p.frequencia} de R$ {p.valor_aposta:.2f}, "
                      f"odd {p.odd_media:.2f}, margem {p.margem_casa:.0%}, {p.anos} ano(s), "
                      f"investimento a {p.taxa_investimento:.0%} ao ano, {p.simulacoes} simulações.")
wr[f"A{fimr + 4}"].font = F_NOTA

wb.save(ARQUIVO)
print(f"Planilha criada: {ARQUIVO}")
print("Abra no Excel e veja a aba 'Resumo'. Todas as verificações devem aparecer como OK.")
